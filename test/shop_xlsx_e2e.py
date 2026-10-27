# Касса Шамсуло в браузере: «Скачать отчёт за месяц в Excel» (№251). Сервер /magazin подменён.
# Скачанный файл открываем openpyxl (как Excel): листы, заголовки, числа, кириллица и спецсимволы.
import subprocess, time, json, os, tempfile
from playwright.sync_api import sync_playwright
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
srv = subprocess.Popen(['python3', '-m', 'http.server', '8132'], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL); time.sleep(1)
SHOP = 'kshamsulotestkod0000000000000'; calls = []
REP = {"month": "2026-10", "revenue": 100, "profit": 20, "owner": 0, "debt": 30, "byDay": {}, "checks": [], "low": [], "expenses": 7, "expBy": {"Маруф": 7}, "expList": [{"id": "e1", "date": "2026-10-25", "who": "Маруф", "note": "обед", "sum": 7}],
       "payBy": {}, "payToday": {}, "sdachi": [{"id": "s1", "date": "2026-10-25", "who": "Маруф", "fact": 50, "exp": 52, "diff": -2, "by": {}, "ok": ""}], "kassa": [{"id": "k1", "date": "2026-10-25", "who": "Маруф", "kind": "out", "sum": 20, "note": "забрал"}],
       "postDebts": [{"name": "Саид", "sum": 65, "since": "2026-10-20"}], "postPays": [{"id": "p1", "date": "2026-10-26", "who": "Саид", "pay": "perevod", "sum": 10}], "itogiHour": 21,
       "days": {"2026-10-24": {"sold": 40, "cashIn": 40, "debtNew": 0, "profit": 8, "purchase": 100, "salary": 0, "pers": 0}, "2026-10-25": {"sold": 60.5, "cashIn": 30.5, "debtNew": 30, "profit": 12, "purchase": 0, "salary": 0, "pers": 7}}, "salary": [], "payList": []}
EX = {"month": "2026-10", "sales": [
        {"date": "2026-10-24", "chek": "bmvaaaaaaaachek01", "kod": "015", "name": "Гвозди <100> & \"Ко\"", "qty": 2, "price": 20, "sum": 40, "zakup": 16, "pay": "nal", "parts": "", "who": "", "seller": "Маруф"},
        {"date": "2026-10-25", "chek": "bmvaaaaaaaachek02", "kod": "TRUBA", "name": "Труба ПП 20", "qty": 3.25, "price": 10, "sum": 32.5, "zakup": 8, "pay": "mix", "parts": "nal 20, alif 12.5", "who": "", "seller": ""},
        {"date": "2026-10-25", "chek": "bmvaaaaaaaachek01r123456", "kod": "015", "name": "Гвозди <100> & \"Ко\"", "qty": -1, "price": 20, "sum": -20, "zakup": 16, "pay": "nal", "parts": "", "who": "", "seller": "Маруф"}],
      "receipts": [{"date": "2026-10-24", "kod": "015", "name": "Гвозди", "qty": 10, "zakup": 6.5, "sum": 65, "post": "Саид", "pay": "dolg", "who": "Али", "fix": False}],
      "debtPays": [{"date": "2026-10-26", "who": "Карим", "sum": 5}]}
def resp(extra=None):
    r = {"ok": True, "role": "owner", "products": [], "debts": [{"name": "Карим", "sum": 30, "since": "2026-10-25", "owner": False}], "receipts": [], "report": REP}
    r.update(extra or {}); return r
def h(route):
    u = route.request.url
    if u.endswith('/magazin'):
        b = json.loads(route.request.post_data); calls.append(b)
        if b.get('action') == 'ping': return route.fulfill(json={"ok": False, "error": "неизвестное действие"})
        if b.get('action') == 'export': return route.fulfill(json=resp({"export": EX}))
        return route.fulfill(json=resp())
    route.fulfill(json={"ok": True, "staff": [], "sessions": [], "chats": []})
def check(c, m):
    print(('ok - ' if c else 'FAIL - ') + m)
    if not c: raise SystemExit(1)
try:
    with sync_playwright() as p:
        br = p.chromium.launch(executable_path=os.environ.get('CHROME') or None); pg = br.new_page(viewport={'width': 390, 'height': 800}, accept_downloads=True); errs = []; pg.on('pageerror', lambda e: errs.append(str(e)))
        pg.route('**/webhook/**', h)
        pg.goto('http://localhost:8132/window.html'); pg.wait_for_timeout(500)
        pg.evaluate("c => { localStorage.setItem('ew_cab', c); document.querySelectorAll('dialog[open]').forEach(d => d.close()); }", SHOP); pg.reload(); pg.wait_for_timeout(1200)
        pg.evaluate("document.querySelectorAll('dialog[open]').forEach(d=>d.close())")
        pg.click('[data-home="fin"]'); pg.wait_for_timeout(500); pg.click('#sp-body [data-fin="shop"]'); pg.wait_for_timeout(800)
        pg.click('[data-shop="tab"][data-k="otchet"]'); pg.wait_for_timeout(400)
        check(pg.locator('[data-shop="xlsx"]').count() == 1, 'в «Отчёте» кнопка «Скачать отчёт за месяц в Excel»')
        with pg.expect_download() as dl: pg.click('[data-shop="xlsx"]')
        d = dl.value; path = os.path.join(tempfile.mkdtemp(), d.suggested_filename); d.save_as(path)
        check(d.suggested_filename == 'otchet-2026-10.xlsx', 'имя файла: %s' % d.suggested_filename)
        ec = [c for c in calls if c.get('action') == 'export']
        check(len(ec) == 1 and ec[0]['month'] == '2026-10', 'запрос выгрузки за месяц отчёта')
        wb = openpyxl.load_workbook(path)
        check(wb.sheetnames == ['По дням', 'Продажи', 'Приём товара', 'Долги', 'Касса'], 'листы: %s' % wb.sheetnames)
        ws = wb['По дням']; rows = list(ws.iter_rows(values_only=True))
        check(rows[0][:3] == ('Дата', 'Продано', 'Пришло денег') and ws['A1'].font.b, 'заголовок жирный')
        check(rows[1] == ('2026-10-24', 40, 40, 0, 8, 100, 0, 0) and rows[3][:3] == ('Итого', 100.5, 70.5), 'по дням и итог: %s' % (rows[3],))
        ws = wb['Продажи']; rows = list(ws.iter_rows(values_only=True))
        check(rows[1][3] == 'Гвозди <100> & "Ко"' and rows[1][4] == 2 and rows[1][10] == 'Маруф', 'спецсимволы и кириллица в названии')
        check(rows[2][4] == 3.25 and 'смешанная (наличные 20, Алиф 12.5)' == rows[2][8] and rows[2][10] == 'руководитель', 'дробное количество, смешанная оплата, продавал руководитель')
        check(rows[3][4] == -1 and 'возврат' in rows[3][8], 'возврат с минусом')
        rows = list(wb['Приём товара'].iter_rows(values_only=True))
        check(rows[1][:8] == ('2026-10-24', '015', 'Гвозди', 10, 6.5, 65, 'Саид', 'в долг'), 'приём: поставщик и долг')
        dv = [r for r in wb['Долги'].iter_rows(values_only=True)]
        flat = [c for r in dv for c in r if c is not None]
        check('Карим' in flat and 30 in flat and 'Саид' in flat and 65 in flat and -10 in flat and 5 in flat, 'долги: нам должны, мы должны, оплаты')
        kv = [c for r in wb['Касса'].iter_rows(values_only=True) for c in r if c is not None]
        check(-2 in kv and -20 in kv and 'обед' in kv, 'касса: сдача, изъятие, личный расход')
        check('Файл отчёта скачан' in pg.inner_text('#sp-body'), 'сообщение о скачивании')
        check(not errs, 'нет ошибок на странице: %s' % errs)
        print('SHOP XLSX E2E PASSED')
finally:
    srv.terminate()
